#!/usr/bin/env python3
"""
Transformer un export CSV de lister_restaurants.py en fichier de prospection Excel.

    python3 prospection_xlsx.py toulouse-peripherie.csv prospection.xlsx --titre "Toulouse et périphérie"

Nécessite openpyxl (pip install openpyxl). Le classeur produit contient :
- Prospection : une ligne par établissement, colonnes de suivi à remplir
  (statut en liste déroulante, dates, interlocuteur, commentaires), liens carte ;
- Synthèse : compteurs par statut, par commune et par type de cuisine (formules) ;
- Aide : source, légende, exemple de ligne remplie.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TRADUCTIONS = {
    "french": "Française", "regional": "Régionale / terroir", "italian": "Italienne", "pizza": "Pizza",
    "italian_pizza": "Pizza", "japanese": "Japonaise", "sushi": "Sushi", "asian": "Asiatique",
    "vietnamese": "Vietnamienne", "chinese": "Chinoise", "china": "Chinoise", "sichuan": "Sichuanaise",
    "thai": "Thaïlandaise", "indian": "Indienne", "lebanese": "Libanaise", "fine_dining": "Gastronomique",
    "gastronomic": "Gastronomique", "burger": "Burger", "steak_house": "Grill / viandes", "steak": "Grill / viandes",
    "grill": "Grill / viandes", "seafood": "Fruits de mer", "fish": "Poisson", "tapas": "Tapas",
    "spanish": "Espagnole", "mexican": "Mexicaine", "tacos": "Tacos", "moroccan": "Marocaine",
    "african": "Africaine", "senegal": "Sénégalaise", "korean": "Coréenne", "crepe": "Crêperie",
    "pancake": "Crêperie", "savory_pancakes": "Crêperie", "breton": "Bretonne", "ramen": "Ramen",
    "noodle": "Nouilles", "noodles": "Nouilles", "poke": "Poké", "hawaiian": "Hawaïenne", "salad": "Salades",
    "sandwich": "Sandwicherie", "turkish": "Turque", "greek": "Grecque", "portuguese": "Portugaise",
    "brazilian": "Brésilienne", "argentinian": "Argentine", "peruvian": "Péruvienne", "chilean": "Chilienne",
    "venezuelan": "Vénézuélienne", "south_american": "Sud-américaine", "latin_american": "Latino-américaine",
    "caribbean": "Antillaise", "antillaise": "Antillaise", "creole": "Créole", "mascarene": "Réunionnaise",
    "mediterranean": "Méditerranéenne", "oriental": "Orientale", "arab": "Orientale", "algerian": "Algérienne",
    "tunisian": "Tunisienne", "couscous": "Couscous", "ethiopian": "Éthiopienne", "armenian": "Arménienne",
    "afghan": "Afghane", "persian": "Persane", "iranienne": "Persane", "nepalese": "Népalaise",
    "tibetan": "Tibétaine", "indonesian": "Indonésienne", "american": "Américaine", "german": "Allemande",
    "alsatian": "Alsacienne", "norwegian": "Norvégienne", "russian": "Russe", "basque": "Basque",
    "cassoulet": "Cassoulet", "traditional": "Traditionnelle", "bistro": "Bistrot", "brasserie": "Brasserie",
    "international": "Internationale", "local": "Produits locaux", "familiale": "Familiale",
    "chicken": "Poulet", "barbecue": "Barbecue", "bbq": "Barbecue", "fondue": "Fondue",
    "hotpot": "Fondue chinoise", "dimsum": "Dim sum", "ravioli": "Raviolis", "pasta": "Pâtes",
    "cafeteria": "Cafétéria", "tea": "Salon de thé", "teahouse": "Salon de thé", "cake": "Pâtisserie",
    "pastry": "Pâtisserie", "breakfast": "Petit-déjeuner", "brunch": "Brunch", "cocktails": "Cocktails",
    "coffee_shop": "Coffee shop", "wok": "Wok", "buffet": "Buffet", "soup": "Soupes", "tartines": "Tartines",
    "smoothie": "Smoothies", "waffle": "Gaufres", "hot_dog": "Hot-dog", "fish_and_chips": "Fish & chips",
    "paleo": "Paléo", "friture": "Friture", "beef_bowl": "Bols de bœuf", "diner": "Diner américain",
    "izakaya": "Izakaya", "delivery": "Livraison", "livraison": "Livraison", "curry": "Curry",
    "kebab": "Kebab", "vegetarian": "Végétarienne", "vegan": "Végane", "bagel": "Bagels",
    "vacant": "Vacant (fermé ?)", "restaurant": "", "bar": "",
}
STATUTS = ["À contacter", "Contacté", "RDV pris", "Devis envoyé", "Client", "Pas intéressé", "À relancer"]
POLICE = "Arial"
BLEU_FONCE, JAUNE_CLAIR, GRIS = "1F4E78", "FFF2CC", "F2F2F2"
BORDURE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def police(**kw) -> Font:
    return Font(name=POLICE, size=kw.pop("size", 10), **kw)


def traduire_cuisine(brut: str) -> str:
    vus, res = set(), []
    for c in [x.strip() for x in brut.split(",") if x.strip()]:
        t = TRADUCTIONS.get(c, c.replace("_", " ").capitalize())
        if t and t not in vus:
            vus.add(t)
            res.append(t)
    return ", ".join(res)


def normaliser_tel(t: str) -> str:
    out = []
    for p in [x.strip() for x in re.split(r"[;/]", t) if x.strip()]:
        d = re.sub(r"[^\d+]", "", p)
        if d.startswith("+33"):
            d = "0" + d[3:]
        out.append(" ".join(d[i:i + 2] for i in range(0, 10, 2)) if len(d) == 10 and d.isdigit() else p)
    return " / ".join(out)


def lire_csv(chemin: str) -> tuple[list[dict], int]:
    """Lit l'export, traduit et dédoublonne (même commune, même nom et même téléphone ou adresse)."""
    with open(chemin, encoding="utf-8", newline="") as f:
        brutes = list(csv.DictReader(f))
    vus: set = set()
    fiches, doublons = [], 0
    for r in brutes:
        nom = re.sub(r"\s{2,}", " ", r["nom"]).strip()
        tel = normaliser_tel(r.get("telephone", ""))
        adresse = r.get("adresse", "").strip()
        ville = r["ville"]
        cles = set()
        if tel:
            cles.add((ville, nom.lower(), "tel", tel))
        if adresse:
            cles.add((ville, nom.lower(), "adr", adresse.lower()))
        if cles & vus:
            doublons += 1
            continue
        vus |= cles
        fiches.append({
            "nom": nom, "cuisine": traduire_cuisine(r.get("cuisine", "")), "adresse": adresse,
            "cp": r.get("code_postal", ""), "ville": ville, "tel": tel, "web": r.get("site_web", "").strip(),
            "horaires": r.get("horaires", ""), "lat": r.get("latitude", ""), "lon": r.get("longitude", ""),
            "osm": r.get("osm_id", ""),
        })
    return fiches, doublons


def entete_ligne(ws, ligne: int, titres: list[str], largeurs: list[int] | None = None) -> None:
    for i, t in enumerate(titres, start=1):
        c = ws.cell(row=ligne, column=i, value=t)
        c.font = police(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=BLEU_FONCE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDURE
        if largeurs:
            ws.column_dimensions[get_column_letter(i)].width = largeurs[i - 1]


def construire(fiches: list[dict], doublons: int, sortie: str, titre: str, source: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospection"
    entetes = ["N°", "Restaurant", "Type de cuisine", "Adresse", "Code postal", "Commune", "Téléphone",
               "Site web", "Horaires", "Statut", "Date 1er contact", "Interlocuteur", "Prochaine action",
               "Date de relance", "Commentaires", "Carte", "Fiche OSM"]
    largeurs = [6, 36, 26, 32, 10, 22, 18, 30, 34, 15, 15, 22, 26, 15, 40, 14, 12]
    entete_ligne(ws, 1, entetes, largeurs)
    ws.row_dimensions[1].height = 30
    COL_STATUT, COL_DATE1, COL_DATE2 = 10, 11, 14
    for r, f in enumerate(fiches, start=2):
        if f["lat"] and f["lon"]:
            carte = f"https://www.google.com/maps/search/?api=1&query={f['lat']},{f['lon']}"
        else:
            carte = "https://www.google.com/maps/search/?api=1&query=" + \
                    re.sub(r"\s+", "+", f"{f['nom']} {f['adresse']} {f['ville']}")
        osm = f"https://www.openstreetmap.org/{f['osm']}" if f["osm"] else ""
        valeurs = [r - 1, f["nom"], f["cuisine"], f["adresse"], f["cp"], f["ville"], f["tel"],
                   f["web"] or None, f["horaires"], STATUTS[0], None, None, None, None, None,
                   "Voir la carte", "Ouvrir" if osm else None]
        for i, v in enumerate(valeurs, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = police()
            c.border = BORDURE
            c.alignment = Alignment(vertical="top", wrap_text=(i in (4, 9, 15)))
            if COL_STATUT <= i <= 15:
                c.fill = PatternFill("solid", start_color=JAUNE_CLAIR)
            if i in (COL_DATE1, COL_DATE2):
                c.number_format = "DD/MM/YYYY"
        for col, url in ((8, f["web"]), (16, carte), (17, osm)):
            if url:
                c = ws.cell(row=r, column=col)
                c.hyperlink = url
                c.font = police(color="0563C1", underline="single")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
    derniere = len(fiches) + 1
    L = get_column_letter
    plage = lambda col: f"Prospection!${col}$2:${col}${derniere}"

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True)
    dv.errorTitle, dv.error = "Statut invalide", "Choisissez un statut dans la liste déroulante."
    ws.add_data_validation(dv)
    dv.add(f"{L(COL_STATUT)}2:{L(COL_STATUT)}{derniere}")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
    dv_date.errorTitle, dv_date.error = "Date invalide", "Saisissez une date au format JJ/MM/AAAA."
    ws.add_data_validation(dv_date)
    dv_date.add(f"{L(COL_DATE1)}2:{L(COL_DATE1)}{derniere}")
    dv_date.add(f"{L(COL_DATE2)}2:{L(COL_DATE2)}{derniere}")
    couleurs = {"Client": ("C6EFCE", "006100"), "Pas intéressé": ("D9D9D9", "595959"),
                "RDV pris": ("FCE4D6", "833C0B"), "Devis envoyé": ("FCE4D6", "833C0B"),
                "À relancer": ("FFEB9C", "9C5700"), "Contacté": ("DDEBF7", "1F4E78")}
    for statut, (fond, texte) in couleurs.items():
        ws.conditional_formatting.add(f"{L(COL_STATUT)}2:{L(COL_STATUT)}{derniere}", CellIsRule(
            operator="equal", formula=[f'"{statut}"'],
            fill=PatternFill("solid", start_color=fond, end_color=fond),
            font=Font(name=POLICE, size=10, bold=True, color=texte)))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{L(len(entetes))}{derniere}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # ------------------------------------------------------------ Synthèse
    sy = wb.create_sheet("Synthèse")
    for col, w in zip("ABCDE", (40, 14, 12, 12, 14)):
        sy.column_dimensions[col].width = w
    sy["A1"] = f"Synthèse de la prospection - {titre}"
    sy["A1"].font = police(bold=True, size=14, color=BLEU_FONCE)
    sy["A2"] = "Les compteurs se mettent à jour automatiquement à partir de la feuille Prospection."
    sy["A2"].font = police(italic=True, color="595959")
    col_statut, col_ville, col_cuisine = L(COL_STATUT), L(6), L(3)

    def pct(r: int, num: str) -> None:
        c = sy.cell(row=r, column=3, value=f"=IF($B$12=0,0,{num}/$B$12)")
        c.font = police()
        c.number_format = "0.0%"
        c.border = BORDURE

    entete_ligne(sy, 4, ["Statut", "Nombre", "Part"])
    for k, statut in enumerate(STATUTS):
        r = 5 + k
        sy.cell(row=r, column=1, value=statut).font = police()
        sy.cell(row=r, column=2, value=f"=COUNTIF({plage(col_statut)},A{r})").font = police()
        pct(r, f"B{r}")
        for c in (1, 2):
            sy.cell(row=r, column=c).border = BORDURE
    r_total = 5 + len(STATUTS)  # 12
    sy.cell(row=r_total, column=1, value="Total des fiches").font = police(bold=True)
    sy.cell(row=r_total, column=2, value=f"=SUM(B5:B{r_total - 1})").font = police(bold=True)
    pct(r_total, f"B{r_total}")
    for c in (1, 2, 3):
        sy.cell(row=r_total, column=c).border = BORDURE
        sy.cell(row=r_total, column=c).fill = PatternFill("solid", start_color=GRIS)
    sy["A14"], sy["B14"] = "Fiches traitées (statut différent de « À contacter »)", f"=B{r_total}-B5"
    sy["A15"], sy["B15"] = "Taux de transformation (clients / fiches traitées)", "=IF(B14=0,0,B9/B14)"
    sy["B15"].number_format = "0.0%"
    sy["A16"], sy["B16"] = "Fiches en attente (RDV, devis, relance)", "=B7+B8+B11"
    for r in (14, 15, 16):
        sy.cell(row=r, column=1).font = police()
        sy.cell(row=r, column=2).font = police(bold=True)

    entete_ligne(sy, 18, ["Données disponibles", "Nombre", "Part"])
    for k, (lib, col) in enumerate([("Fiches avec téléphone", L(7)), ("Fiches avec adresse", L(4)),
                                    ("Fiches avec site web", L(8)), ("Fiches avec horaires", L(9)),
                                    ("Fiches avec type de cuisine", col_cuisine)]):
        r = 19 + k
        sy.cell(row=r, column=1, value=lib).font = police()
        sy.cell(row=r, column=2, value=f"=COUNTA({plage(col)})").font = police()
        pct(r, f"B{r}")
        for c in (1, 2):
            sy.cell(row=r, column=c).border = BORDURE

    communes = list(dict.fromkeys(f["ville"] for f in fiches))
    r0 = 26
    entete_ligne(sy, r0, ["Commune", "Fiches", "Part", "Traitées", "Clients"])
    for k, ville in enumerate(communes):
        r = r0 + 1 + k
        sy.cell(row=r, column=1, value=ville).font = police()
        sy.cell(row=r, column=2, value=f"=COUNTIF({plage(col_ville)},A{r})").font = police()
        pct(r, f"B{r}")
        sy.cell(row=r, column=4, value=f'=B{r}-COUNTIFS({plage(col_ville)},A{r},{plage(col_statut)},"{STATUTS[0]}")').font = police()
        sy.cell(row=r, column=5, value=f'=COUNTIFS({plage(col_ville)},A{r},{plage(col_statut)},"Client")').font = police()
        for c in (1, 2, 4, 5):
            sy.cell(row=r, column=c).border = BORDURE
    r_fin_communes = r0 + len(communes)
    r = r_fin_communes + 1
    sy.cell(row=r, column=1, value="Total").font = police(bold=True)
    for c, col in ((2, "B"), (4, "D"), (5, "E")):
        sy.cell(row=r, column=c, value=f"=SUM({col}{r0 + 1}:{col}{r_fin_communes})").font = police(bold=True)
    for c in (1, 2, 3, 4, 5):
        sy.cell(row=r, column=c).border = BORDURE
        sy.cell(row=r, column=c).fill = PatternFill("solid", start_color=GRIS)

    compteur = Counter(t for f in fiches for t in [x.strip() for x in f["cuisine"].split(",")] if t)
    r0c = r + 3
    entete_ligne(sy, r0c, ["Type de cuisine (20 plus fréquents)", "Fiches", "Part"])
    for k, (cuisine, _) in enumerate(compteur.most_common(20)):
        r = r0c + 1 + k
        sy.cell(row=r, column=1, value=cuisine).font = police()
        sy.cell(row=r, column=2, value=f'=COUNTIF({plage(col_cuisine)},"*"&A{r}&"*")').font = police()
        pct(r, f"B{r}")
        for c in (1, 2):
            sy.cell(row=r, column=c).border = BORDURE
    sy.cell(row=r0c + 22, column=1,
            value="Une fiche avec plusieurs types de cuisine est comptée dans chacun d'eux.").font = police(italic=True, color="595959")
    sy.freeze_panes = "A4"

    # ------------------------------------------------------------ Aide
    ai = wb.create_sheet("Aide")
    ai.column_dimensions["A"].width, ai.column_dimensions["B"].width = 30, 95
    ai["A1"] = "Mode d'emploi du fichier de prospection"
    ai["A1"].font = police(bold=True, size=14, color=BLEU_FONCE)
    textes = [
        ("Objet", f"Suivi de la prospection commerciale des restaurants : {titre}. Une ligne par établissement, "
                  "avec les colonnes de suivi à compléter au fil des contacts."),
        ("Source des données", f"{source} {len(fiches) + doublons} établissements listés, {doublons} doublons retirés "
                               f"(même commune, même nom et même téléphone ou adresse), {len(fiches)} fiches conservées, "
                               f"{len(communes)} communes."),
        ("Fiabilité", "Les données OpenStreetMap sont contributives : certains établissements n'ont pas de téléphone ou "
                      "d'adresse, d'autres ont pu fermer. Vérifiez la fiche (liens « Voir la carte » et « Ouvrir ») "
                      "avant tout démarchage."),
        ("Colonnes A à I", "Données de l'établissement : numéro, nom, type de cuisine (traduit des catégories OpenStreetMap), "
                           "adresse, code postal, commune, téléphone au format français, site web (cliquable), horaires "
                           "(notation OpenStreetMap : Mo-Fr 12:00-14:00 = du lundi au vendredi de 12 h à 14 h). "
                           "Vous pouvez les corriger si besoin."),
        ("Colonnes J à O (fond jaune)", "Colonnes de suivi à remplir : Statut (liste déroulante), Date 1er contact, "
                                        "Interlocuteur, Prochaine action, Date de relance, Commentaires. "
                                        "Les dates se saisissent au format JJ/MM/AAAA."),
        ("Colonnes P et Q", "Liens automatiques : position exacte sur Google Maps, et fiche OpenStreetMap de l'établissement."),
        ("Statuts", " · ".join(STATUTS) + ". Le statut colore la cellule ; la feuille Synthèse compte les fiches "
                                          "par statut, par commune (avec fiches traitées et clients) et par type de cuisine."),
        ("Filtres", "Les en-têtes de la feuille Prospection ont des filtres : filtrez par commune ou par type de cuisine, "
                    "triez par date de relance pour préparer vos appels."),
        ("Mettre à jour", "Régénérer le CSV : python3 lister_restaurants.py --peripherie --format csv --sortie export.csv, "
                          "puis ce classeur : python3 prospection_xlsx.py export.csv prospection.xlsx"),
    ]
    for k, (t, texte) in enumerate(textes):
        r = 3 + k
        ai.cell(row=r, column=1, value=t).font = police(bold=True)
        ai.cell(row=r, column=1).alignment = Alignment(vertical="top")
        c = ai.cell(row=r, column=2, value=texte)
        c.font = police()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ai.row_dimensions[r].height = 45 if len(texte) > 110 else 30
    r_ex = 3 + len(textes) + 1
    ai.cell(row=r_ex, column=1, value="Exemple de ligne remplie").font = police(bold=True, size=12, color=BLEU_FONCE)
    exemple = ["Chez Exemple (fictif)", "Française", "1 Rue de l'Exemple", "31000", "Toulouse", "05 61 00 00 00",
               "https://exemple.fr", "Mo-Sa 12:00-14:00,19:00-22:00", "RDV pris", "02/10/2026",
               "Mme Martin, gérante", "Présenter l'offre en rendez-vous", "09/10/2026",
               "Intéressée par une démonstration, préfère être appelée le matin"]
    for k, (h, v) in enumerate(zip(entetes[1:15], exemple)):
        r = r_ex + 1 + k
        c1 = ai.cell(row=r, column=1, value=h)
        c1.font, c1.border = police(bold=True), BORDURE
        c2 = ai.cell(row=r, column=2, value=v)
        c2.font, c2.border = police(), BORDURE
        if k >= 8:
            c2.fill = PatternFill("solid", start_color=JAUNE_CLAIR)
    ai.cell(row=r_ex + len(exemple) + 2, column=1,
            value="Ligne d'exemple fictive : elle n'est pas dans la feuille Prospection.").font = police(italic=True, color="595959")

    wb.calculation.fullCalcOnLoad = True
    wb.save(sortie)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("csv", help="export CSV de lister_restaurants.py")
    p.add_argument("xlsx", help="classeur Excel à produire")
    p.add_argument("--titre", default="Restaurants", help="titre de la synthèse (ex. « Toulouse et périphérie »)")
    p.add_argument("--source", default=None, help="phrase décrivant la source, pour la feuille Aide")
    args = p.parse_args()
    fiches, doublons = lire_csv(args.csv)
    source = args.source or (f"Liste extraite d'OpenStreetMap (API Overpass) le {date.today():%d/%m/%Y} "
                             "avec le script lister_restaurants.py.")
    construire(fiches, doublons, args.xlsx, args.titre, source)
    print(f"{len(fiches)} fiches ({doublons} doublons retirés) écrites dans {args.xlsx}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
